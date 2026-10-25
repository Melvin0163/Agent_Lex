"""Loading user documents (PDF, DOCX, DOC, TXT, MD, HTML, RTF) and splitting them into clauses."""
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

SUPPORTED = (".pdf", ".docx", ".doc", ".txt", ".md", ".html", ".htm", ".rtf")


def load_document(path) -> str:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"File not found: {path}")
    ext = p.suffix.lower()

    if ext in (".txt", ".md"):
        return p.read_text(encoding="utf-8", errors="ignore")

    if ext in (".html", ".htm"):
        raw = p.read_text(encoding="utf-8", errors="ignore")
        raw = re.sub(r"(?is)<(script|style).*?</\1>", " ", raw)
        raw = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</li>|</h\d>", "\n", raw)
        return re.sub(r"<[^>]+>", " ", raw)

    if ext == ".rtf":
        raw = p.read_text(encoding="utf-8", errors="ignore")
        raw = re.sub(r"\\par[d]?", "\n", raw)
        raw = re.sub(r"\\[a-z]+-?\d* ?|[{}]", "", raw)
        return raw

    if ext == ".pdf":
        from pypdf import PdfReader
        text = "\n".join((pg.extract_text() or "") for pg in PdfReader(str(p)).pages)
        if len(text.strip()) < 50:
            raise ValueError("No readable text in this PDF (it is probably a scanned image). "
                             "Please upload a text-based PDF, DOCX or TXT.")
        return text

    if ext == ".docx":
        from docx import Document
        d = Document(str(p))
        parts = [para.text for para in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(x for x in parts if x.strip())

    if ext == ".doc":
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            raise ValueError("Old .doc files need LibreOffice to convert. "
                             "Easier: open it in Word, Save As .docx, then upload again.")
        out = tempfile.mkdtemp()
        subprocess.run([soffice, "--headless", "--convert-to", "docx", "--outdir", out, str(p)],
                       check=True, capture_output=True)
        return load_document(Path(out) / (p.stem + ".docx"))

    raise ValueError(f"Unsupported file type: {ext}. Use one of: {', '.join(SUPPORTED)}")


def chunk_text(text, size=900):
    """Group paragraphs into chunks of roughly `size` characters (used by the RAG pipeline)."""
    paras = [x.strip() for x in text.split("\n") if x.strip()]
    chunks, cur = [], ""
    for para in paras:
        if cur and len(cur) + len(para) > size:
            chunks.append(cur)
            cur = ""
        cur = (cur + "\n" + para).strip()
        while len(cur) > size * 2:
            chunks.append(cur[:size * 2])
            cur = cur[size * 2:]
    if cur:
        chunks.append(cur)
    return chunks


_HEADING = re.compile(
    r"^\s*(?:(?:clause|section|article)\s+)?(\d{1,2}(?:\.\d{1,2})?)[.)]?\s+([A-Z][^\n]{0,80})$", re.I)


_SECTION_LINE = re.compile(
    r"^\s*(?:section|clause|article)\s*[-\u2013\u2014:.]?\s*(\d{1,2}|[IVXL]{1,6})\b\s*[-\u2013\u2014:.]?\s*(.{0,60})$", re.I)


def _split_by_sections(lines, text):
    preamble, clauses, cur, i = [], [], None, 0
    while i < len(lines):
        ln = lines[i]
        m = _SECTION_LINE.match(ln)
        if m:
            if cur:
                clauses.append(cur)
            head = m.group(2).strip(" :.-\u2013\u2014")
            if not head:                       # title on the next non-empty line
                j = i + 1
                while j < len(lines) and not lines[j].strip():
                    j += 1
                if j < len(lines) and len(lines[j].strip()) <= 70:
                    head, i = lines[j].strip().rstrip(":."), j
            head = head.title() if head.isupper() else head
            cur = {"no": m.group(1), "heading": head or f"Section {m.group(1)}", "body": ""}
        elif ln.strip():
            if cur is None:
                preamble.append(ln.strip())
            else:
                cur["body"] = (cur["body"] + "\n" + ln.strip()).strip()
        i += 1
    if cur:
        clauses.append(cur)
    for c in clauses:
        if not c["body"]:
            c["body"] = c["heading"]
    return " ".join(preamble).strip(), clauses


def split_clauses(text):
    """Split a contract/policy into clauses: [{"no", "heading", "body"}].

    Recognises numbered headings ("1. Probation", "Clause 3 - Bond"). Falls back to
    paragraph chunks when the document has no numbering.
    """
    lines = [ln.rstrip() for ln in text.splitlines()]
    # Documents with "SECTION 3" / "Section - 1" headings (title often on the next line): those sections are
    # the clauses, and numbered lines inside them (2.1, 4.2 ...) are sub-clauses that stay in the body.
    # Used only when sections are clause-sized; a long handbook's chapters are too big to be clauses.
    if sum(bool(_SECTION_LINE.match(ln)) for ln in lines) >= 2:
        pre, secs = _split_by_sections(lines, text)
        if len(secs) >= 2 and sum(len(s["body"]) for s in secs) / len(secs) <= 4000:
            return pre, secs
    preamble, clauses, cur = [], [], None
    for ln in lines:
        m = _HEADING.match(ln)
        if m and len(ln) < 100:
            if cur:
                clauses.append(cur)
            rest = m.group(2).strip()
            # "1. Probation is 3 months..." (heading + body on one line) vs "1. Probation"
            if len(rest.split()) > 8:
                head, body = " ".join(rest.split()[:5]) + "...", rest
            else:
                head, body = rest.rstrip(":.- "), ""
            cur = {"no": m.group(1), "heading": head, "body": body}
        elif cur is not None:
            if ln.strip():
                cur["body"] = (cur["body"] + "\n" + ln.strip()).strip()
        elif ln.strip():
            preamble.append(ln.strip())
    if cur:
        clauses.append(cur)

    if len(clauses) < 2:   # unnumbered document: use paragraph chunks
        clauses = [{"no": str(i), "heading": c.split("\n")[0][:60], "body": c}
                   for i, c in enumerate(chunk_text(text, 700), 1)]
        preamble = []
    for c in clauses:
        if not c["body"]:
            c["body"] = c["heading"]
    return " ".join(preamble).strip(), clauses
